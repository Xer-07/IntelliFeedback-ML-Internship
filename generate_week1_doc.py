import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets inner padding for table cells (in dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout(doc, text, title="KEY STRATEGIC TAKEAWAY"):
    """Adds a stylish callout box with a colored left accent border."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="1A365D"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"📌 {title}: ")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    run_body = p.add_run(text)
    run_body.font.name = "Calibri"
    run_body.font.size = Pt(10.5)
    run_body.font.italic = True
    run_body.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table(table, col_widths, headers, data, alt_bg="F7FAFC", header_bg="1A365D"):
    """Formats a modern table with colored header, padding, and alternating row backgrounds."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    # Data Rows
    for r_idx, row_data in enumerate(data):
        row_cells = table.add_row().cells
        bg_color = alt_bg if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            row_cells[c_idx].width = col_widths[c_idx]
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=140, right=140)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # Apply subtle border
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E0"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="1A365D"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_document():
    doc = docx.Document()
    
    # Page Setup: Standard Letter, 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Set default style font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # DOCUMENT HEADER / BANNER
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_kicker = title_p.add_run("Yuvaintern Virtual ML Internship | Week 1 Project Deliverable\n")
    run_kicker.font.size = Pt(9.5)
    run_kicker.font.bold = True
    run_kicker.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    
    run_title = title_p.add_run("Project Planning & Strategy Development")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(14)
    run_sub = subtitle_p.add_run("Comprehensive Machine Learning Charter: IntelliFeedback – Customer Review Sentiment & Feedback Intelligence Engine")
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    # Metadata Table
    meta_table = doc.add_table(rows=1, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_headers = ["Internship Program", "Track", "Deliverable Type", "Duration"]
    meta_values = [["YuvaIntern 5-Week Virtual Internship", "Machine Learning (ML)", "Week 1 Project Charter & Strategy", "5 Weeks (Weeks 1 – 5)"]]
    col_widths = [Inches(1.8), Inches(1.5), Inches(2.2), Inches(1.0)]
    format_table(meta_table, col_widths, meta_headers, meta_values, header_bg="2B6CB0")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # SECTION 1: EXECUTIVE SUMMARY
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    r1 = h1.add_run("1. Executive Summary")
    r1.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    r1.font.bold = True

    p = doc.add_paragraph(
        "In modern e-commerce and digital service ecosystems, customer-generated text (reviews, feedback tickets, social commentary) "
        "constitutes one of the richest yet most underutilized streams of operational intelligence. The primary objective of this 5-week "
        "machine learning initiative is to conceptualize, architect, validate, and deploy IntelliFeedback, an end-to-end Natural Language "
        "Processing (NLP) and Machine Learning platform capable of automated multiclass sentiment classification, key aspect extraction, "
        "and customer dissatisfaction signal detection."
    )
    p = doc.add_paragraph(
        "This strategic project plan establishes the technical foundation for the entire project lifecycle. It articulates the core problem "
        "statement, provides an empirical literature review, outlines hypothesis formulations, defines end-to-end data pipelines, sets up "
        "a rigorous validation matrix, schedules a 5-week milestone timeline, and conducts a proactive risk mitigation assessment. "
        "By adhering to enterprise ML engineering principles, the project balances computational efficiency with predictive rigor, "
        "providing a turnkey solution ready for real-world deployment."
    )
    
    add_callout(
        doc,
        "IntelliFeedback bridges the gap between raw unstructured customer text and actionable executive decisions by deploying a dual-tiered "
        "architecture: a high-throughput, low-latency TF-IDF/Linear baseline for bulk stream ingestion coupled with a contextual transformer "
        "(DistilBERT) for nuanced, ambiguous feedback resolution.",
        "PROJECT CHARTER PRINCIPLE"
    )

    # -------------------------------------------------------------
    # SECTION 2: PROBLEM DEFINITION & SIGNIFICANCE
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("2. Problem Definition & Business Significance")
    r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    r.font.bold = True

    h2 = doc.add_heading(level=2)
    r = h2.add_run("2.1 The Core Challenge")
    r.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    
    p = doc.add_paragraph(
        "Digital enterprises process tens of thousands of customer reviews and support interactions on a daily basis. Despite this "
        "abundance of data, companies face a critical 'Feedback Paralysis' bottleneck: human auditing of unstructured text is prohibitively "
        "costly, slow, and subjective. Consequently, critical alerts regarding product defects, logistics failures, and severe customer "
        "dissatisfaction are often surfaced weeks after initial occurrence, resulting in avoidable customer churn and brand erosion."
    )
    p = doc.add_paragraph(
        "Furthermore, superficial star-rating aggregates fail to capture nuanced sentiment: a 4-star review may contain acute complaints "
        "regarding shipping delays ('Great product, but delivery was a nightmare'), whereas a 2-star review might praise customer support "
        "despite a broken component. Without granular, context-aware Natural Language Processing, businesses remain blind to the true root causes."
    )

    h2 = doc.add_heading(level=2)
    r = h2.add_run("2.2 Significance and Potential Impact")
    r.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    p = doc.add_paragraph("Deploying an intelligent, automated feedback analytics pipeline yields direct, quantifiable benefits:")
    
    bullet_items = [
        ("Real-Time Churn Prevention: ", "Immediate triage of severe negative feedback (escalation to customer retention teams within minutes)."),
        ("Operational Cost Reduction: ", "Automating the classification of incoming sentiment reduces manual ticket tagging overhead by over 75%."),
        ("Product Engineering Visibility: ", "Categorization of sentiment by functional themes (e.g., packaging, quality, usability) accelerates sprint triage."),
        ("Brand Reputation Monitoring: ", "Instant detection of abnormal negative sentiment spikes indicates potential outages or supply chain disruptions.")
    ]
    for b_title, b_desc in bullet_items:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        rt = bp.add_run(b_title)
        rt.bold = True
        rt.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        rd = bp.add_run(b_desc)

    # -------------------------------------------------------------
    # SECTION 3: RESEARCH & LITERATURE REVIEW
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("3. Research Background & Literature Review")
    r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    r.font.bold = True

    p = doc.add_paragraph(
        "Sentiment analysis and opinion mining have evolved through three distinct paradigms in computational linguistics and machine learning:"
    )

    lit_table = doc.add_table(rows=1, cols=4)
    lit_headers = ["Paradigm", "Core Techniques", "Key Strengths", "Critical Limitations"]
    lit_data = [
        [
            "1. Lexicon-Based\n(Rule-based)",
            "VADER, SentiWordNet, TextBlob dictionaries",
            "Zero training data required; deterministic; ultra-fast execution (<1ms).",
            "Fails on negation ('not bad'), sarcasm, slang, and context-dependent domain jargon."
        ],
        [
            "2. Statistical / Classical ML\n(Feature Engineering)",
            "Bag-of-Words, TF-IDF n-grams + Logistic Regression, SVM, Naive Bayes",
            "High sample efficiency; interpretable feature weights; highly scalable in production.",
            "Ignores word order and contextual semantics; struggles with polysemy and long-range dependencies."
        ],
        [
            "3. Deep Contextual / Transformers\n(Representation Learning)",
            "BERT (Devlin et al., 2018), RoBERTa (Liu et al., 2019), DistilBERT (Sanh et al., 2019)",
            "State-of-the-art accuracy; captures bidirectional context, idioms, syntactic subtleties.",
            "High inference latency; compute-intensive training requirements; susceptibility to fine-tuning instability."
        ]
    ]
    lit_widths = [Inches(1.5), Inches(1.8), Inches(1.8), Inches(1.4)]
    format_table(lit_table, lit_widths, lit_headers, lit_data)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    p = doc.add_paragraph(
        "Industry benchmarks (such as the Amazon Product Reviews benchmark and Yelp Open Dataset) demonstrate that while transformer "
        "models yield the highest overall Macro F1-score, production environments often suffer from prohibitive inference costs. "
        "Therefore, modern industry engineering standards recommend benchmarking robust classical pipelines (TF-IDF + Ridge/Logistic Regression) "
        "alongside lightweight transformer distillation (DistilBERT) to strike the optimal trade-off between predictive quality and system latency."
    )

    # -------------------------------------------------------------
    # SECTION 4: HYPOTHESIS FORMULATION & PROJECT GOALS
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("4. Hypothesis Formulation & Project Goals")
    r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    r.font.bold = True

    h2 = doc.add_heading(level=2)
    r = h2.add_run("4.1 Core Hypotheses")
    r.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    hypotheses = [
        ("Hypothesis 1 (Feature Representation): ", 
         "Sub-linear TF-IDF scaling coupled with bi-gram and tri-gram tokens will outperform standard unigram Bag-of-Words by at least 5% in macro F1-score due to its ability to retain negations (e.g., 'not good', 'hardly works')."),
        ("Hypothesis 2 (Linear vs. Non-Linear Models): ", 
         "Linear Support Vector Machines (LinearSVC) and Logistic Regression with L2 regularization will demonstrate comparable accuracy to tree-based ensembles (Random Forest, XGBoost) on high-dimensional sparse TF-IDF vectors, while requiring less than one-tenth of the training time."),
        ("Hypothesis 3 (Contextual Deep Learning): ", 
         "Fine-tuning a pre-trained DistilBERT architecture will provide a statistically significant uplift (>6% macro F1-score) on subtle and sarcastic reviews where n-gram models inherently produce false positives.")
    ]
    for h_title, h_desc in hypotheses:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        rh = p.add_run(h_title)
        rh.bold = True
        rh.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        p.add_run(h_desc)

    h2 = doc.add_heading(level=2)
    r = h2.add_run("4.2 SMART Project Objectives")
    r.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    obj_table = doc.add_table(rows=1, cols=3)
    obj_headers = ["Dimension", "SMART Objective Specification", "Target Milestone"]
    obj_data = [
        ["Specific", "Build a multiclass NLP classifier categorizing reviews into Positive, Neutral, and Negative classes with confidence scores.", "Week 3 Deliverable"],
        ["Measurable", "Attain an overall Macro F1-Score of ≥ 0.88 and a balanced accuracy of ≥ 85% on held-out test data.", "Week 4 Deliverable"],
        ["Achievable", "Utilize established public e-commerce datasets (Amazon/Yelp, 20,000+ curated records) and open-source Scikit-Learn/HuggingFace libraries.", "Week 2 Deliverable"],
        ["Relevant", "Directly solve feedback routing latency, enabling instant triage of high-risk negative complaints.", "Week 5 Deliverable"],
        ["Time-Bound", "Complete all design, preprocessing, model tuning, evaluation, and interactive demo deployment within 5 calendar weeks.", "Weeks 1 – 5"]
    ]
    obj_widths = [Inches(1.2), Inches(3.8), Inches(1.5)]
    format_table(obj_table, obj_widths, obj_headers, obj_data)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 5: METHODOLOGY & TECHNICAL APPROACH
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("5. Methodology & Technical Approach")
    r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    r.font.bold = True

    p = doc.add_paragraph(
        "The project follows a standard CRISP-DM (Cross-Industry Standard Process for Data Mining) and MLOps lifecycle, "
        "structured into five rigorous pipeline stages:"
    )

    steps = [
        ("Step 1: Data Acquisition & Preprocessing: ",
         "Acquire a benchmark multi-domain customer review dataset (e.g., Amazon Consumer Reviews dataset, Kaggle / Stanford SNAP). "
         "Execute text normalization: HTML decoding, lowercase normalization, contraction expansion ('won't' -> 'will not'), punctuation cleaning "
         "while preserving sentiment emojis, stopword pruning (ensuring negation words like 'no', 'not', 'never' are retained), and WordNet Lemmatization."),
        
        ("Step 2: Exploratory Data Analysis (EDA): ",
         "Analyze class label distribution, review length histograms, word cloud distributions per sentiment category, and n-gram frequency distributions. "
         "Identify anomalies such as duplicate reviews, bots, and empty strings."),
        
        ("Step 3: Feature Extraction & Engineering: ",
         "Construct two parallel representation tracks: (A) Statistical Track using sub-linear TF-IDF vectorization with n-gram range (1, 3), max_features=10,000, "
         "and min_df=5; (B) Deep Learning Track using Hugging Face AutoTokenizer for DistilBERT with dynamic padding and attention masking."),
        
        ("Step 4: Model Exploration & Hyperparameter Tuning: ",
         "Train baseline models: Multinomial Naive Bayes, Logistic Regression (L2 penalty), Support Vector Classifier (LinearSVC), and Random Forest. "
         "Perform GridSearchCV / RandomizedSearchCV for hyperparameter optimization (regularization strength C, n-gram ranges, solver types). "
         "Fine-tune pre-trained DistilBERT with AdamW optimizer, warmup schedules, and learning rate 2e-5."),
        
        ("Step 5: Pipeline Packaging & Inference Interface: ",
         "Serialize the winning preprocessing and model artifacts using Joblib/ONNX. Develop a modular Python inference class and an interactive "
         "web interface (Streamlit) for real-time text input and batch CSV scoring.")
    ]
    for s_title, s_desc in steps:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        rs = p.add_run(s_title)
        rs.bold = True
        rs.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        p.add_run(s_desc)

    # -------------------------------------------------------------
    # SECTION 6: EVALUATION CRITERIA & VALIDATION STRATEGY
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("6. Evaluation Criteria & Validation Strategy")
    r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    r.font.bold = True

    p = doc.add_paragraph(
        "In customer feedback classification, class distributions are characteristically imbalanced (positive reviews typically outnumber "
        "negative reviews 3:1). Consequently, relying solely on raw Accuracy is misleading. The project adopts a comprehensive evaluation matrix:"
    )

    eval_table = doc.add_table(rows=1, cols=3)
    eval_headers = ["Metric", "Mathematical Definition / Purpose", "Strategic Target"]
    eval_data = [
        ["Macro F1-Score", "Unweighted average of F1 across Positive, Neutral, and Negative classes; penalizes poor performance on minority classes.", "≥ 0.88"],
        ["Negative Class Recall", "Proportion of actual negative reviews correctly flagged: TP_neg / (TP_neg + FN_neg). Essential for churn triage.", "≥ 0.90"],
        ["Balanced Accuracy", "Average recall obtained on each class, guarding against majority class bias.", "≥ 0.85"],
        ["ROC-AUC (One-vs-Rest)", "Separability metric evaluating probability calibration across multiclass thresholds.", "≥ 0.92"],
        ["Inference Latency", "95th percentile execution time per single review prediction in production.", "< 20 ms (classical)\n< 120 ms (DistilBERT)"]
    ]
    eval_widths = [Inches(1.8), Inches(3.4), Inches(1.3)]
    format_table(eval_table, eval_widths, eval_headers, eval_data)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    p = doc.add_paragraph(
        "Validation Strategy: Data will be partitioned into 70% Training, 15% Validation, and 15% Blind Testing. "
        "Model tuning on the training set will utilize Stratified 5-Fold Cross-Validation with fixed random seeds (random_state=42) "
        "to guarantee statistical reproducibility across experimental runs."
    )

    # -------------------------------------------------------------
    # SECTION 7: 5-WEEK TIMELINE & MILESTONES
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("7. 5-Week Implementation Timeline & Milestones")
    r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    r.font.bold = True

    p = doc.add_paragraph(
        "The project execution is mapped to the 5-week structure of the YuvaIntern Machine Learning Virtual Internship, "
        "with clearly demarcated tasks and tangible deliverables for each sprint:"
    )

    time_table = doc.add_table(rows=1, cols=4)
    time_headers = ["Timeline", "Sprint Focus", "Key Technical Tasks", "Expected Deliverables"]
    time_data = [
        [
            "Week 1\n(Current)",
            "Project Planning & Strategy Development",
            "• Problem formulation & business case analysis\n• Literature review & architecture selection\n• GitHub repository structuring & documentation\n• Risk assessment and evaluation framework design",
            "• Week 1 Project Plan Document (.DOCX)\n• Official GitHub Repository Setup\n• Portal submission & progress summary"
        ],
        [
            "Week 2",
            "Data Collection, Ingestion & EDA",
            "• Benchmark dataset acquisition (20k+ reviews)\n• Exploratory Data Analysis & visual profiling\n• Text cleaning & normalization pipeline development\n• Vocabulary analysis, n-grams & class imbalance audit",
            "• Jupyter Notebook (EDA & Data Cleaning)\n• Curated Clean Dataset artifacts\n• Statistical Data Profile Summary"
        ],
        [
            "Week 3",
            "Feature Engineering & Baseline Modeling",
            "• TF-IDF feature extraction (n-gram vectorization)\n• Word embeddings & lexical polarity scoring\n• Baseline classification (Naive Bayes, Logistic Regression, LinearSVC)\n• Initial benchmark performance reporting",
            "• Feature extraction module (Python src)\n• Model comparison benchmark table\n• Baseline evaluation report"
        ],
        [
            "Week 4",
            "Model Tuning, Deep Learning & Optimization",
            "• Hyperparameter optimization via GridSearchCV\n• Deep Learning exploration (DistilBERT fine-tuning)\n• Error analysis on misclassified instances & edge cases\n• Threshold calibration and model persistence",
            "• Finalized Model Artifacts (.joblib / .pt)\n• Confusion matrices & ROC curves\n• In-depth Model Evaluation Report"
        ],
        [
            "Week 5",
            "Packaging, Deployment & Final Showcase",
            "• Inference API & Streamlit web application build\n• End-to-end testing with unseen customer text\n• Comprehensive final project documentation & video demo\n• Final internship report submission",
            "• Working Web Application (Streamlit)\n• Complete GitHub codebase with CI/Docs\n• Final Internship Completion Report"
        ]
    ]
    time_widths = [Inches(1.1), Inches(1.6), Inches(2.3), Inches(1.5)]
    format_table(time_table, time_widths, time_headers, time_data)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 8: RISK ASSESSMENT & MITIGATION PLAN
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("8. Comprehensive Risk Assessment & Mitigation Plan")
    r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    r.font.bold = True

    p = doc.add_paragraph(
        "A rigorous risk assessment was conducted to anticipate technical, statistical, and infrastructural bottlenecks. "
        "The risk matrix below outlines each identified vulnerability, its severity profile, and the active contingency plan:"
    )

    risk_table = doc.add_table(rows=1, cols=5)
    risk_headers = ["Risk Factor", "Likelihood", "Impact", "Risk Severity", "Proactive Mitigation Strategy"]
    risk_data = [
        [
            "Severe Class Imbalance",
            "High",
            "High",
            "Critical",
            "Adopt class-weighted loss functions (e.g., class_weight='balanced' in Scikit-Learn), SMOTE / text data augmentation (EDA techniques), and evaluate on Macro F1 rather than Accuracy."
        ],
        [
            "Linguistic Ambiguity & Sarcasm",
            "Medium",
            "Medium",
            "Moderate",
            "Incorporate bi-gram and tri-gram n-grams to preserve negation tokens. Integrate transformer attention (DistilBERT) which captures bidirectional context."
        ],
        [
            "Overfitting on Domain Lexicons",
            "Medium",
            "High",
            "High",
            "Apply L1/L2 regularization on linear models; prune rare words (min_df=5) and domain-specific stopwords; validate using Stratified 5-Fold Cross Validation."
        ],
        [
            "Compute Constraints for Transformers",
            "High",
            "Low",
            "Moderate",
            "Use quantized / distilled architectures (DistilBERT or MobileBERT) instead of full-size BERT. Leverage cloud GPU (Google Colab / Kaggle T4) for fine-tuning."
        ],
        [
            "Timeline Slippage Across Sprints",
            "Low",
            "High",
            "Moderate",
            "Adopt decoupled modular code architecture (separate data, features, modeling, and UI modules) with pre-planned milestone buffers."
        ]
    ]
    risk_widths = [Inches(1.4), Inches(0.9), Inches(0.8), Inches(0.9), Inches(2.5)]
    format_table(risk_table, risk_widths, risk_headers, risk_data)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 9: TECHNICAL STACK & ENVIRONMENT
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("9. Technical Stack & Tooling Architecture")
    r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    r.font.bold = True

    stack_table = doc.add_table(rows=1, cols=3)
    stack_headers = ["Layer / Category", "Tools & Frameworks", "Role in Project Lifecycle"]
    stack_data = [
        ["Programming Language", "Python 3.10+", "Primary runtime for ML development and scripting."],
        ["Data Processing & EDA", "Pandas, NumPy, Matplotlib, Seaborn", "Tabular data manipulation, statistical profiling, and distribution visualization."],
        ["NLP & Preprocessing", "NLTK, spaCy, RegEx, scikit-learn", "Tokenization, lemmatization, stopword cleaning, TF-IDF vectorization."],
        ["Machine Learning Models", "Scikit-Learn, LightGBM, Hugging Face Transformers", "Classification algorithms, pipeline orchestration, DistilBERT fine-tuning."],
        ["Evaluation & Tuning", "Scikit-Learn Metrics, Yellowbrick", "Cross-validation, confusion matrices, ROC curves, calibration curves."],
        ["Application & UI", "Streamlit, Joblib", "Interactive live demo web application and model serialization."],
        ["Version Control", "Git, GitHub", "Source code management, sprint deliverables tracking, reproducible code sharing."]
    ]
    stack_widths = [Inches(1.8), Inches(2.2), Inches(2.5)]
    format_table(stack_table, stack_widths, stack_headers, stack_data)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 10: REFERENCES & BIBLIOGRAPHY
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("10. References & Empirical Literature")
    r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    r.font.bold = True

    refs = [
        "1. Devlin, J., Chang, M. W., Lee, K., & Toutanova, K. (2018). BERT: Pre-training of Deep Bidirectional Transformers for Language Understanding. arXiv preprint arXiv:1810.04805.",
        "2. Sanh, V., Debut, L., Chaumond, J., & Wolf, T. (2019). DistilBERT, a distilled version of BERT: smaller, faster, cheaper and lighter. arXiv preprint arXiv:1910.01108.",
        "3. Pang, B., & Lee, L. (2008). Opinion mining and sentiment analysis. Foundations and Trends in Information Retrieval, 2(1–2), 1-135.",
        "4. Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830.",
        "5. McAuley, J., Targett, C., Shi, Q., & van den Hengel, A. (2015). Image-based recommendations on styles and substitutes. In Proceedings of the 38th International ACM SIGIR Conference on Research and Development in Information Retrieval."
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r_ref = p.add_run(ref)
        r_ref.font.size = Pt(9.5)
        r_ref.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    # Save document
    output_dir = "docs"
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "Week1_Project_Planning_and_Strategy_Report.docx")
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    build_document()
